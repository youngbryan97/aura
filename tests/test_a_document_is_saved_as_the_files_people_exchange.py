"""A document made in a program she built is saved as files other programs open, and opened back from them.

The formats are code in the frame (core/rebuilding/document_formats.js), so a
part only calls them; checked here with the programs that read those files.
"""
from __future__ import annotations

import asyncio
import base64
import io
import shutil

import pytest

from core.rebuilding.the_program_as_built import Part, ProgramAsBuilt

_WORK = """
const page = app.make("div", {class: "page", contenteditable: "true", "aria-label": "Document"});
page.innerHTML = '<h1>Quarterly Letter</h1><p>Dear <b>Ms. Reyes</b>, thank you for <i>everything</i>.</p>'
  + '<p style="text-align:center">Café — naïve “quotes”</p><ul><li>First point</li><li>Second point</li></ul>'
  + '<table><tr><td>Region</td><td>Sales</td></tr><tr><td>North</td><td>42</td></tr></table>';
app.work.append(page); app.doc = page;
"""


async def _written(tmp_path, kind: str) -> bytes:
    from playwright.async_api import async_playwright

    page_file = ProgramAsBuilt("Writer", parts=[Part("work area", _WORK)]).write(tmp_path / "w.html")
    async with async_playwright() as pw:
        browser = await pw.chromium.launch()
        page = await browser.new_page()
        await page.goto(page_file.resolve().as_uri())
        made = await page.evaluate("""async (kind) => { const bytes = new Uint8Array(await app.formats.write(kind).arrayBuffer());
            let s = ""; for (const b of bytes) s += String.fromCharCode(b); return btoa(s); }""", kind)
        await browser.close()
    return base64.b64decode(made)


@pytest.mark.asyncio
async def test_a_docx_opens_in_a_reader_of_word_files(tmp_path):
    import docx

    document = docx.Document(io.BytesIO(await _written(tmp_path, "docx")))
    paragraphs = document.paragraphs
    assert paragraphs[0].text == "Quarterly Letter" and paragraphs[0].style.name.lower().startswith("heading 1")
    letter = next(p for p in paragraphs if p.text.startswith("Dear"))
    assert letter.text == "Dear Ms. Reyes, thank you for everything."
    assert any(r.bold and r.text == "Ms. Reyes" for r in letter.runs) and any(r.italic and r.text == "everything" for r in letter.runs)
    assert any("Café — naïve “quotes”" in p.text for p in paragraphs)
    assert [p.text for p in paragraphs if "point" in p.text] == ["First point", "Second point"]
    assert document.tables[0].cell(1, 1).text == "42"


@pytest.mark.asyncio
@pytest.mark.skipif(shutil.which("textutil") is None, reason="textutil is the macOS reader of these files")
@pytest.mark.parametrize("kind", ["docx", "rtf", "odt", "html"])
async def test_the_system_reads_each_kind(tmp_path, kind):
    path = tmp_path / f"letter.{kind}"
    path.write_bytes(await _written(tmp_path, kind))
    read = await asyncio.create_subprocess_exec("textutil", "-convert", "txt", "-stdout", str(path), stdout=asyncio.subprocess.PIPE)
    text = (await read.communicate())[0].decode("utf-8", "ignore")
    assert "Quarterly Letter" in text and "Ms. Reyes" in text and "Café" in text, text[:300]


@pytest.mark.asyncio
async def test_plain_kinds_say_what_the_document_says(tmp_path):
    md = (await _written(tmp_path, "md")).decode()
    assert md.startswith("# Quarterly Letter") and "**Ms. Reyes**" in md and "- First point" in md
    txt = (await _written(tmp_path, "txt")).decode()
    assert "Dear Ms. Reyes, thank you for everything." in txt and "• First point" in txt


@pytest.mark.asyncio
async def test_a_docx_from_another_program_is_opened_back(tmp_path):
    """Read back from a compressed .docx written by another program, not only from its own."""
    import docx
    from playwright.async_api import async_playwright

    made = docx.Document()
    made.add_heading("From elsewhere", level=1)
    para = made.add_paragraph("Plain then ")
    para.add_run("bold").bold = True
    held = io.BytesIO()
    made.save(held)
    data_url = "data:application/octet-stream;base64," + base64.b64encode(held.getvalue()).decode()
    page_file = ProgramAsBuilt("Writer", parts=[Part("work area", _WORK)]).write(tmp_path / "r.html")
    async with async_playwright() as pw:
        browser = await pw.chromium.launch()
        page = await browser.new_page()
        await page.goto(page_file.resolve().as_uri())
        html = await page.evaluate("(url) => app.formats.read({name: 'other.docx', dataUrl: url})", data_url)
        rtf = await page.evaluate("""() => app.formats.read({name: 'x.rtf', text: '{\\\\rtf1\\\\ansi{\\\\fonttbl{\\\\f0 Arial;}}Hello \\\\b world\\\\b0\\\\par Next}'})""")
        await browser.close()
    assert "<h1>From elsewhere</h1>" in html and "<strong>bold</strong>" in html
    assert "Hello world" in rtf and "Next" in rtf and "Arial" not in rtf


def test_the_formats_a_program_saves_are_read_from_what_is_written():
    from core.rebuilding.what_the_frame_gives import formats_named

    article = ("Microsoft Word is a word processor. Its native file format is DOCX. Word can also save documents as "
               "Rich Text Format and plain text, and open OpenDocument files. The company sells Word Documents online.")
    assert formats_named([article]) == ["docx", "odt", "rtf", "txt"]
    assert formats_named(["A drawing program with layers."]) == []
    assert formats_named(["and export it to my Desktop as a .docx"]) == ["docx"]


@pytest.mark.asyncio
async def test_a_format_named_is_given_by_the_frame_with_nothing_asked_of_her_model(tmp_path):
    from core.rebuilding.rebuilding_a_program import rebuild
    from tests.test_a_program_is_rebuilt_by_checking_it import _NoCorpus, _Script

    script = _Script()
    done = await rebuild("Writer", script, tmp_path, corpus=_NoCorpus(), online=False, asked="rebuild it, and let me save my letters as .docx files")
    saving = {o.feature.name: o for o in done.built.outcomes}["Save as Word Document (.docx)"]
    assert saving.kept and saving.held == 1
    assert not [p for p in script.asked if 'part for the feature "Save as Word Document' in p]
    assert 'app.formats.save("docx")' in done.built.path.read_text()


@pytest.mark.asyncio
async def test_asking_for_a_format_added_to_a_build_is_done_by_the_frame(tmp_path):
    """"Add .docx export to the word processor you built": the change is read from the words and given by code."""
    from core.rebuilding.changing_what_was_built import change_it
    from core.rebuilding.rebuilding_a_program import rebuild
    from tests.test_a_program_is_rebuilt_by_checking_it import _NoCorpus, _Script

    done = await rebuild("Writer", _Script(), tmp_path, corpus=_NoCorpus(), online=False)

    async def says_nothing_useful(prompt, schema, max_tokens):
        return None

    changed = await change_it(done.built.path.parent, "add an export to .docx files to it", says_nothing_useful)
    assert [o.kept for o in changed.outcomes] == [True], changed.summary()
    assert 'app.formats.save("docx")' in (done.built.path.parent / "index.html").read_text()


@pytest.mark.asyncio
async def test_a_check_gives_a_real_file_of_the_kind_it_names(tmp_path):
    """LIVE 2026-10-05 a check gave plain text named test.docx, and a right Open was left out for it."""
    from core.rebuilding.checks_a_person_makes import Check, run_checks

    opening = Part("Open", 'app.command({label: "Open", menu: "File", run: () => app.formats.open()});', ["Open"])
    check = Check.model_validate({"feature": "Open", "steps": [{"do": "type", "value": "Existing content"}, {"do": "click", "target": "Open"},
                                                                {"do": "give_file", "target": "test.docx", "value": "New document content"}],
                                  "expect": [{"see": "text", "target": "New document content"}, {"see": "no_text", "target": "Existing content"}]})
    runs = await run_checks(ProgramAsBuilt("W", parts=[Part("work area", _WORK), opening]).write(tmp_path / "o.html"), [check])
    assert runs[0].held, runs[0].why


@pytest.mark.asyncio
async def test_a_new_build_never_takes_up_the_leavings_of_another(tmp_path):
    from core.rebuilding.rebuilding_a_program import rebuild
    from tests.test_a_program_is_rebuilt_by_checking_it import _NoCorpus, _Script

    stale = tmp_path / "writer"
    stale.mkdir()
    (stale / "checks.json").write_text("[]")
    done = await rebuild("Writer", _Script(), tmp_path, corpus=_NoCorpus(), online=False)
    assert done.built.path.parent.name == "writer-2" and any(o.kept for o in done.built.outcomes)


@pytest.mark.asyncio
async def test_opening_and_saving_a_document_are_the_frames(tmp_path):
    """LIVE 2026-10-05 her model's Open read a .docx as plain text; opening and saving are the frame's, by code."""
    from core.rebuilding.checks_a_person_makes import run_checks
    from core.rebuilding.what_a_program_does import Feature
    from core.rebuilding.what_the_frame_gives import what_the_frame_gives

    mine = [Feature(name="Open a document", how="File, Open, pick a .docx or .txt", shows="it loads"),
            Feature(name="Save document", how="File, Save", shows="a file is saved"),
            Feature(name="Bold", how="select, press B", shows="bold")]
    features, given = what_the_frame_gives(mine, [], "rebuild it; I keep my letters as .docx files")
    assert set(given) >= {"Open a document", "Save document"} and "Bold" not in given
    parts = [Part("work area", _WORK), *(g.part for g in given.values())]
    runs = await run_checks(ProgramAsBuilt("W", parts=parts).write(tmp_path / "p.html"), [c for g in given.values() for c in g.checks])
    assert all(r.held for r in runs), [r.why for r in runs if not r.held]
